from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "docs" / "screenshots"
OUT = ROOT / "docs" / "汇报.pptx"

PAPER = RGBColor(0xF7, 0xF4, 0xEB)
INK = RGBColor(0x1F, 0x2A, 0x24)
MUTED = RGBColor(0x6C, 0x75, 0x6F)
GREEN = RGBColor(0x2F, 0x5D, 0x4E)
GREEN_DARK = RGBColor(0x22, 0x43, 0x39)
RUST = RGBColor(0xB4, 0x69, 0x3F)
LINE = RGBColor(0xD8, 0xD1, 0xC0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xEC, 0xE7, 0xDA)

FONT = "Microsoft YaHei"
MONO = "Consolas"

W, H = 13.333, 7.5


def solid(shape, color: RGBColor, line: RGBColor | None = None, width=1.0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(width)
    return shape


def rect(slide, x, y, w, h, color, radius=False, line=None, width=1.0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    solid(shape, color, line, width)
    if radius:
        shape.adjustments[0] = 0.06
    return shape


def text(
    slide,
    x,
    y,
    w,
    h,
    value,
    *,
    size=16,
    bold=False,
    color=INK,
    font=FONT,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    spacing=1.15,
    margins=(0.04, 0.03, 0.04, 0.03),
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = Inches(margins[0])
    frame.margin_top = Inches(margins[1])
    frame.margin_right = Inches(margins[2])
    frame.margin_bottom = Inches(margins[3])
    lines = value.split("\n")
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = spacing
        run = paragraph.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def bullets(slide, x, y, w, h, items, *, size=16, color=INK, gap=8, bullet_color=GREEN):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    for index, item in enumerate(items):
        level = 0
        if isinstance(item, tuple):
            item, level = item
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.level = level
        paragraph.line_spacing = 1.18
        paragraph.space_after = Pt(gap)
        run = paragraph.add_run()
        run.text = ("• " if level == 0 else "– ") + item
        run.font.name = FONT
        run.font.size = Pt(size - (2 if level else 0))
        run.font.color.rgb = bullet_color if level == 0 else MUTED
    return box


def chip(slide, x, y, w, label, *, fill=SOFT, color=GREEN, size=11):
    shape = rect(slide, x, y, w, 0.32, fill, radius=True)
    text(slide, x, y + 0.015, w, 0.29, label, size=size, color=color,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return shape


def header(slide, kicker, title, page):
    rect(slide, 0, 0, W, H, PAPER)
    text(slide, 0.72, 0.42, 9.8, 0.28, kicker, size=11, bold=True, color=RUST)
    text(slide, 0.72, 0.70, 11.9, 0.62, title, size=27, bold=True, color=INK)
    rect(slide, 0.72, 1.36, 11.9, 0.018, LINE)
    text(slide, 11.65, 7.06, 0.9, 0.25, f"{page:02d}", size=10, color=MUTED,
         align=PP_ALIGN.RIGHT)
    text(slide, 0.72, 7.06, 6.8, 0.25,
         "个人学习空间 · 《数据库系统概论》大作业", size=9, color=MUTED)


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


def picture_fit(slide, filename, x, y, w, h, *, frame=True):
    path = SHOTS / filename
    if not path.exists():
        raise FileNotFoundError(path)
    with Image.open(path) as image:
        iw, ih = image.size
    ratio = min(w / iw, h / ih)
    pw, ph = iw * ratio, ih * ratio
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    if frame:
        rect(slide, x - 0.05, y - 0.05, w + 0.1, h + 0.1, WHITE, radius=True, line=LINE, width=0.8)
    slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    return px, py, pw, ph


def code_block(slide, x, y, w, h, code, *, size=10):
    rect(slide, x, y, w, h, GREEN_DARK, radius=True)
    text(slide, x + 0.16, y + 0.12, w - 0.32, h - 0.24, code, size=size,
         color=RGBColor(0xEC, 0xF2, 0xEA), font=MONO, spacing=1.1)


def table(slide, x, y, w, h, rows, widths, *, font_size=10.5, header_size=10.5):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    for column, width in enumerate(widths):
        tbl.columns[column].width = Inches(width)
    for row_index, row in enumerate(rows):
        tbl.rows[row_index].height = Inches(h / len(rows))
        for column, value in enumerate(row):
            cell = tbl.cell(row_index, column)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = GREEN if row_index == 0 else (
                WHITE if row_index % 2 else SOFT
            )
            frame = cell.text_frame
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.LEFT
            run = paragraph.add_run()
            run.text = str(value)
            run.font.name = FONT
            run.font.size = Pt(header_size if row_index == 0 else font_size)
            run.font.bold = row_index == 0
            run.font.color.rgb = WHITE if row_index == 0 else INK
    return tbl


def card(slide, x, y, w, h, title, body, *, number=None, accent=GREEN):
    rect(slide, x, y, w, h, WHITE, radius=True, line=LINE, width=0.7)
    rect(slide, x, y, 0.08, h, accent, radius=False)
    top = y + 0.18
    if number is not None:
        text(slide, x + 0.24, top, 0.7, 0.42, number, size=25, bold=True, color=accent)
        text(slide, x + 0.95, top + 0.02, w - 1.15, 0.34, title, size=15, bold=True, color=INK)
        text(slide, x + 0.24, top + 0.52, w - 0.46, h - 0.7, body, size=12,
             color=MUTED, spacing=1.1)
    else:
        text(slide, x + 0.24, top, w - 0.46, 0.34, title, size=15, bold=True, color=accent)
        text(slide, x + 0.24, top + 0.42, w - 0.46, h - 0.58, body, size=12,
             color=MUTED, spacing=1.12)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, GREEN_DARK)
    rect(slide, 8.9, -1.2, 6.2, 6.2, RGBColor(0x2B, 0x50, 0x44), radius=True)
    rect(slide, 10.6, 3.4, 4.6, 4.6, RGBColor(0x37, 0x5E, 0x50), radius=True)
    text(slide, 0.85, 0.82, 7.8, 0.3, "数据库系统概论 · 大作业汇报", size=12, bold=True,
         color=RGBColor(0xD9, 0xC8, 0xA8))
    text(slide, 0.85, 1.55, 8.6, 1.25, "个人学习空间", size=46, bold=True, color=WHITE)
    text(slide, 0.85, 2.85, 7.4, 1.0,
         "把与大模型的聊天记录\n变成可查阅、可批注、可复习的课程资料",
         size=22, color=RGBColor(0xE8, 0xE3, 0xD3), spacing=1.22)
    chip(slide, 0.85, 4.55, 1.85, "PostgreSQL 16")
    chip(slide, 2.82, 4.55, 1.75, "FastAPI SSR")
    chip(slide, 4.69, 4.55, 1.45, "Text2SQL")
    text(slide, 0.85, 6.35, 7.6, 0.32,
         "选题 2 · 模拟数据演示 · 汇报人：____________", size=12,
         color=RGBColor(0xC9, 0xD3, 0xCB))
    notes(slide, "开场 20 秒：介绍选题、两个界面形态（传统网站 + Text2SQL）以及模拟数据声明。")


def slide_why(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "01 / 问题与目标", "问答很零散，复习时却找不到上下文", 2)

    text(slide, 0.72, 1.62, 5.2, 0.35, "三个真实痛点", size=13, bold=True, color=GREEN)
    card(slide, 0.72, 2.05, 3.7, 1.55, "记录分散", "Qoder、ChatGPT、DeepSeek 各有一套导出格式，\n历史问答难以统一查找。", number="01")
    card(slide, 0.72, 3.78, 3.7, 1.55, "缺少上下文", "只搜到一句话，却不知道它属于哪门课、\n哪次追问、前后文是什么。", number="02", accent=RUST)
    card(slide, 0.72, 5.51, 3.7, 1.28, "复习低效", "看过的结论没有批注、没有到期提醒，\n无法形成间隔复习。", number="03", accent=GREEN)

    text(slide, 5.05, 1.62, 6.8, 0.35, "系统目标", size=13, bold=True, color=GREEN)
    goals = [
        ("导入", "统一 JSON / Markdown 格式，解析、查重、事务入库"),
        ("组织", "课程 + 来源 + 标签三维归类，支持课程共享"),
        ("复习", "消息级批注、复习队列、1→2→4→8 天间隔"),
        ("洞察", "课程 / 月度 / 标签 / 批注统计与导出"),
    ]
    for index, (name, body) in enumerate(goals):
        y = 2.05 + index * 1.07
        rect(slide, 5.05, y, 7.25, 0.9, WHITE, radius=True, line=LINE, width=0.7)
        text(slide, 5.24, y + 0.12, 1.0, 0.32, name, size=15, bold=True, color=GREEN)
        text(slide, 6.18, y + 0.12, 5.9, 0.62, body, size=12, color=MUTED)

    rect(slide, 5.05, 6.48, 7.25, 0.42, SOFT, radius=True)
    text(slide, 5.22, 6.54, 6.95, 0.3,
         "设计原则：以数据库为中心，手写参数化 SQL，不用 ORM 隐藏关系模型。",
         size=11, bold=True, color=GREEN)
    notes(slide, "30 秒：从三个痛点切入，说明系统不是聊天工具，而是把聊天记录结构化为学习资料。")


def slide_scope(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "02 / 角色与范围", "两类系统角色，三种课程角色", 3)

    cards = [
        ("学生 alice", "5 门课程 owner\n3 个平台来源\n30 个会话 / 420 条消息", GREEN),
        ("同学 bob", "操作系统 editor\n数据库 viewer\n可共享、可复习，不可越权", RUST),
        ("管理员 admin", "用户启停\n系统聚合统计\n不读取他人正文与批注", GREEN_DARK),
    ]
    for index, (name, body, accent) in enumerate(cards):
        x = 0.72 + index * 4.05
        card(slide, x, 1.62, 3.7, 1.72, name, body, accent=accent)

    text(slide, 0.72, 3.66, 4.0, 0.35, "十项功能需求", size=13, bold=True, color=GREEN)
    rows = [
        ["ID", "能力", "关键数据库点"],
        ["R01", "记录导入", "JSONB 预览、唯一键、单事务"],
        ["R02", "课程与标签", "M:N、部分唯一索引、触发器"],
        ["R03", "检索阅读", "trgm + ILIKE、全局搜索、上下文"],
        ["R04", "批注复习", "1:N、复合唯一键、事件流水"],
        ["R05", "统计导出", "JOIN、FILTER、窗口函数、CTE"],
        ["R06", "用户权限", "应用层 + SQL 作用域 + DB 角色"],
        ["R07", "Text2SQL", "只读视图、SQL 守卫、审计"],
        ["R08", "管理", "聚合查询、账号启停"],
        ["R09", "快速采集", "令牌哈希、限流、Bearer 复用事务"],
        ["R10", "离线学习", "IndexedDB、事件 ID 幂等同步"],
    ]
    table(slide, 0.72, 4.05, 7.45, 2.55, rows, [0.72, 1.55, 5.18], font_size=9.5, header_size=10)
    picture_fit(slide, "01-dashboard.png", 8.55, 3.78, 4.0, 2.85)
    notes(slide, "30 秒：强调 alice/bob 的课程级权限与 admin 的聚合边界；R01–R10 与报告需求表对应。")


def slide_data(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "03 / 数据来源", "没有真实导出，就用可复现的模拟数据", 4)

    rect(slide, 0.72, 1.62, 6.35, 2.25, GREEN_DARK, radius=True)
    text(slide, 0.95, 1.78, 5.9, 0.3, "统一导入格式", size=13, bold=True,
         color=RGBColor(0xD9, 0xC8, 0xA8))
    code_block(slide, 0.95, 2.15, 5.9, 1.53,
               '{"platform": "qoder",\n'
               ' "conversations": [{\n'
               '   "external_id": "qoder-db-001",\n'
               '   "title": "SQL 执行顺序", "course_code": "DB",\n'
               '   "tags": ["SQL"], "messages": [...] }]}', size=9)

    stats = [
        ("5", "门课程"), ("30", "个会话"), ("420", "条消息"),
        ("66", "条批注"), ("59", "个标签"), ("5", "个导入批次"),
    ]
    for index, (value, label) in enumerate(stats):
        x = 7.35 + (index % 3) * 1.86
        y = 1.62 + (index // 3) * 1.12
        rect(slide, x, y, 1.68, 0.96, WHITE, radius=True, line=LINE, width=0.7)
        text(slide, x + 0.12, y + 0.08, 1.44, 0.45, value, size=24, bold=True, color=GREEN)
        text(slide, x + 0.12, y + 0.58, 1.44, 0.25, label, size=10, color=MUTED)

    cards = [
        ("JSON 为主", "结构化字段完整；已适配 ChatGPT conversations.json 与 Qoder 会话导出。"),
        ("Markdown 为辅", "# 标题 + 课程/标签元信息 + 用户/助手分段；兼容 DeepSeek 原生标题。"),
        ("可复现", "make_samples.py 生成样例，seed_demo.py 走真实导入服务。"),
    ]
    for index, (name, body) in enumerate(cards):
        x = 0.72 + index * 4.06
        card(slide, x, 4.15, 3.78, 1.42, name, body, accent=RUST if index == 1 else GREEN)

    rect(slide, 0.72, 5.82, 12.0, 0.72, RGBColor(0xF2, 0xE4, 0xD8), radius=True)
    text(slide, 0.95, 5.94, 11.55, 0.5,
         "数据声明：所有聊天内容均为围绕课程知识点生成的模拟数据，不含真实个人信息；报告、页面与导出文件均明确标注。",
         size=12, bold=True, color=RUST)
    notes(slide, "25 秒：主动声明模拟数据；JSON/Markdown 两条导入路径，样例可重新生成。")


def slide_er(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "04 / 概念设计", "8 个核心实体，12 组联系", 5)

    bullets(slide, 0.72, 1.62, 5.35, 4.7, [
        "用户、课程、来源、会话、消息、标签、批注、复习项",
        "课程成员、导入批次、Text2SQL 日志、复习事件、采集令牌是操作型实体",
        "批注与复习项按弱实体处理，依赖消息 / 批注存在",
        "用户 M:N 课程（带角色）；会话 M:N 标签",
        "用户 1:N 来源 → 会话 → 消息 → 批注",
        "同一批注可被不同用户加入各自复习队列",
    ], size=15, gap=12)
    rect(slide, 0.72, 5.95, 5.35, 0.66, SOFT, radius=True)
    text(slide, 0.9, 6.04, 5.0, 0.48,
         "删除策略：消息随会话级联；会话随课程置空；\n有会话的来源禁止删除。", size=11, color=GREEN)
    picture_fit(slide, "16-er-chen.png", 6.25, 1.5, 6.35, 5.35)
    notes(slide, "35 秒：指着 E-R 图讲 8 个实体与 12 组联系；强调批注/复习弱实体和删除策略。")


def slide_schema(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "05 / 逻辑设计", "14 张表：核心业务与操作型数据分离", 6)

    rows = [
        ["领域", "表", "关键约束"],
        ["账户与权限", "users · course_members", "系统角色；每课一个 owner；复合主键"],
        ["内容", "sources · conversations · messages", "(source, external_id)、(conversation, seq_no) 唯一"],
        ["组织", "courses · tags · conversation_tags", "(created_by, code, semester) 唯一；标签同 owner"],
        ["学习", "annotations · review_items · review_logs", "UNIQUE(user_id, annotation_id)；事件流水幂等"],
        ["采集与运维", "import_batches · api_tokens · text2sql_log", "预览 token 唯一；令牌只存哈希；执行审计"],
    ]
    table(slide, 0.72, 1.62, 11.9, 3.3, rows, [1.7, 4.2, 6.0], font_size=11.5, header_size=11.5)

    bullets(slide, 0.72, 5.22, 5.8, 1.5, [
        "单列代理键 + 关联表，避免复合业务键扩散",
        "会话不冗余 owner，消息不冗余课程 / 来源",
        "不存 message_count 等可派生计数",
    ], size=12, gap=5)
    bullets(slide, 6.85, 5.22, 5.75, 1.5, [
        "保留 payload JSONB 作为预览期受控冗余",
        "row_count 是 Text2SQL 审计快照，不是统计缓存",
        "count(*) 实时计算，保证统计与明细一致",
    ], size=12, gap=5)
    notes(slide, "35 秒：按五组讲表；重点说明为什么没有冗余 owner_id 和计数列。")


def slide_norm(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "06 / 范式与完整性", "先满足范式，再只保留有理由的冗余", 7)

    left = [
        ("1NF", "所有列原子；标签、成员、会话标签拆关联表，不存逗号分隔串。"),
        ("2NF", "核心表单列代理键；关联表非键属性依赖整个复合键。"),
        ("3NF / BCNF", "会话不冗余 owner；消息不冗余课程与来源；避免传递依赖。"),
        ("受控冗余", "import_batches.payload 仅预览期；row_count 为审计快照。"),
    ]
    for index, (name, body) in enumerate(left):
        y = 1.62 + index * 1.13
        rect(slide, 0.72, y, 5.95, 0.97, WHITE, radius=True, line=LINE, width=0.7)
        text(slide, 0.94, y + 0.12, 1.55, 0.32, name, size=14, bold=True, color=GREEN)
        text(slide, 2.42, y + 0.11, 4.05, 0.75, body, size=11.5, color=MUTED, spacing=1.1)

    text(slide, 7.0, 1.62, 5.5, 0.35, "完整性约束", size=13, bold=True, color=GREEN)
    rows = [
        ["类型", "例子"],
        ["主键 / 外键", "14 张表全部有 PK；关键关系均有 FK"],
        ["UNIQUE", "会话外部 ID、消息序号、标签名"],
        ["CHECK", "角色枚举、消息非空、结束时间 ≥ 开始"],
        ["部分唯一", "每门课最多一个 owner"],
        ["触发器", "标签与会话同 owner；仅剩一个 owner 时拒绝移除"],
        ["删除行为", "CASCADE / SET NULL / RESTRICT 按语义选择"],
    ]
    table(slide, 7.0, 2.02, 5.6, 2.6, rows, [1.45, 4.15], font_size=10, header_size=10)

    code_block(slide, 7.0, 4.9, 5.6, 1.42,
               "-- 不用生成列：避免 content → hash 传递依赖\n"
               "CREATE INDEX idx_messages_content_md5\n"
               "  ON app.messages (md5(content));\n"
               "-- 中文子串检索\n"
               "CREATE INDEX ... USING gin (content gin_trgm_ops);", size=9)
    notes(slide, "35 秒：1NF/2NF/3NF 各一句；重点讲受控冗余与表达式索引的取舍。")


def slide_import(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "07 / 导入流水线", "解析、预览、确认三段式，确认阶段全有或全无", 8)

    steps = [
        ("1", "上传", "JSON / MD\n≤ 5 MB"),
        ("2", "解析校验", "字段、角色\n时间、上限"),
        ("3", "持久化预览", "JSONB + token\n30 分钟"),
        ("4", "勾选映射", "课程 / 标签\n查重状态"),
        ("5", "事务入库", "advisory lock\nON CONFLICT"),
    ]
    for index, (number, name, body) in enumerate(steps):
        x = 0.72 + index * 2.47
        rect(slide, x, 1.62, 2.15, 1.48, WHITE, radius=True, line=LINE, width=0.7)
        text(slide, x + 0.16, 1.73, 0.4, 0.35, number, size=18, bold=True, color=RUST)
        text(slide, x + 0.58, 1.75, 1.4, 0.3, name, size=13, bold=True, color=INK)
        text(slide, x + 0.16, 2.2, 1.85, 0.7, body, size=10.5, color=MUTED, spacing=1.15)
        if index < len(steps) - 1:
            text(slide, x + 2.12, 2.05, 0.36, 0.4, "→", size=15, bold=True, color=GREEN)

    picture_fit(slide, "08-import-preview.png", 0.72, 3.3, 7.45, 3.3)
    text(slide, 8.5, 3.35, 4.1, 0.35, "演示文件", size=13, bold=True, color=GREEN)
    rows = [
        ["文件", "预期"],
        ["qoder_db.json", "6 会话 / 84 消息，首次入库"],
        ["duplicate_db.json", "全部重复，新增 0"],
        ["broken.json", "字段级错误，直接拒绝"],
    ]
    table(slide, 8.5, 3.74, 4.1, 1.45, rows, [1.55, 2.55], font_size=9.5, header_size=9.5)
    bullets(slide, 8.5, 5.38, 4.15, 1.25, [
        "失败整体回滚，批次单独标记 failed",
        "重复导入只补新增消息",
        "平台导出与 Bearer 接口共用同一事务",
    ], size=11.5, gap=5)
    notes(slide, "40 秒：三段式导入；强调预览不写业务表、确认阶段单事务与幂等。")


def slide_search(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "08 / 检索与阅读", "中文子串检索 + 会话内上下文", 9)

    bullets(slide, 0.72, 1.62, 5.3, 3.8, [
        "标题与消息正文联合匹配，课程 / 标签 / 日期筛选",
        "pg_trgm GIN + ILIKE ESCAPE，参数化查询",
        "count(*) OVER() 一次查询拿到分页总数",
        "SQL 侧截取命中片段，避免传输整段长回答",
        "打开会话按 seq_no 排序，保留完整上下文",
        "关键词在标题和正文中高亮显示",
    ], size=14, gap=10)
    code_block(slide, 0.72, 5.55, 5.3, 1.2,
               "SELECT ..., count(*) OVER() AS total\n"
               "FROM app.messages m\n"
               "WHERE m.content ILIKE %(kw)s ESCAPE '\\'\n"
               "ORDER BY m.conversation_id, m.seq_no;", size=9)
    picture_fit(slide, "05-conversation-detail.png", 6.3, 1.55, 6.3, 5.2)
    notes(slide, "35 秒：搜索“隔离级别”进入会话详情，讲 trgm、分页和上下文。")


def slide_review(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "09 / 批注与复习", "把“看懂了”变成“下次还会”", 10)

    cards = [
        ("四类批注", "note / question / correction / summary", GREEN),
        ("解决状态", "open → resolved；支持编辑与删除", RUST),
        ("复习队列", "pending / reviewing / mastered", GREEN_DARK),
    ]
    for index, (name, body, accent) in enumerate(cards):
        x = 0.72 + index * 2.55
        card(slide, x, 1.62, 2.35, 1.25, name, body, accent=accent)

    bullets(slide, 0.72, 3.15, 6.0, 2.5, [
        "批注挂在消息上；一条消息可有多条批注",
        "复习项唯一键是 (user_id, annotation_id)，同一批注可被多人复习",
        "到期查询使用部分索引 WHERE status <> 'mastered'",
        "复习一次：间隔 1 → 2 → 4 → 8 天，due_date 顺延",
        "标记掌握后退出热点队列，但保留历史",
    ], size=13, gap=8)
    code_block(slide, 0.72, 5.75, 6.0, 0.98,
               "UPDATE app.review_items\n"
               "SET interval_days = interval_days * 2,\n"
               "    due_date = current_date + interval_days * 2\n"
               "WHERE id = %s AND user_id = %s;", size=8.8)
    picture_fit(slide, "10-reviews.png", 7.0, 1.55, 5.6, 5.25)
    notes(slide, "40 秒：从批注页加入复习，再到队列点“复习一次”，展示间隔翻倍与到期日变化。")


def slide_stats(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "10 / 统计与导出", "把课程里的聚合查询真正用起来", 11)

    snippets = [
        ("课程排名", "RANK() OVER (ORDER BY message_count DESC)"),
        ("解决率", "COUNT(*) FILTER (WHERE status = 'resolved')"),
        ("月度补零", "generate_series(date_trunc('month', ...), ...)"),
        ("重复检测", "GROUP BY md5(content) HAVING count(*) > 1"),
    ]
    for index, (name, sql) in enumerate(snippets):
        y = 1.62 + index * 1.02
        rect(slide, 0.72, y, 5.85, 0.86, WHITE, radius=True, line=LINE, width=0.7)
        text(slide, 0.9, y + 0.1, 1.35, 0.28, name, size=11.5, bold=True, color=GREEN)
        text(slide, 2.25, y + 0.1, 4.15, 0.58, sql, size=10.5, color=MUTED, font=MONO, spacing=1.05)

    bullets(slide, 0.72, 5.85, 5.85, 1.0, [
        "全部计数实时计算，不落冗余统计表",
        "CSV 统计导出 + 单会话 Markdown 导出",
    ], size=11.5, gap=4)
    picture_fit(slide, "11-stats.png", 6.85, 1.55, 5.75, 5.25)
    notes(slide, "35 秒：指统计页的课程排名、月度活跃、标签 Top、重复检测；说明导出。")


def slide_t2s(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "11 / Text2SQL", "自然语言查询，但权限在数据库层兜底", 12)

    flow = ["自然语言", "LLM / 本地规则", "SQL 守卫", "只读视图", "结果 + 审计"]
    for index, label in enumerate(flow):
        x = 0.72 + index * 2.42
        rect(slide, x, 1.62, 2.05, 0.64, GREEN if index in (0, 4) else GREEN_DARK,
             radius=True)
        text(slide, x, 1.7, 2.05, 0.4, label, size=11.5, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER)
        if index < len(flow) - 1:
            text(slide, x + 2.06, 1.72, 0.3, 0.35, "→", size=13, bold=True, color=RUST)

    left = [
        "只允许多一条 SELECT / WITH，拒绝 DDL、DML、注释",
        "对象白名单：13 个 t2s 视图，禁止 app / pg_catalog",
        "包裹 200 行上限；t2s_reader 只读事务、3 秒超时",
        "视图 security_barrier，按 app.user_id 过滤",
        "成功、拒绝、异常全部写 text2sql_log",
        "无 Key 走 mock，配置 Key 走 OpenAI 兼容接口",
    ]
    bullets(slide, 0.72, 2.65, 6.15, 3.4, left, size=12.5, gap=8)
    code_block(slide, 0.72, 5.95, 6.15, 0.82,
               "t2s_reader: CONNECT + USAGE(t2s) + SELECT 视图\n"
               "default_transaction_read_only = on", size=9.2)
    picture_fit(slide, "13-text2sql-result.png", 7.15, 2.55, 5.45, 4.2)
    notes(slide, "45 秒：提交“每个课程有多少会话和消息？”；强调多层守卫和只读角色兜底。")


def slide_permissions(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "12 / 权限模型", "应用层、SQL 作用域、数据库约束三层防护", 13)

    rows = [
        ["操作", "owner", "editor", "viewer", "admin"],
        ["查看课程 / 会话", "✓", "✓", "✓", "仅聚合"],
        ["编辑课程 / 共享", "✓", "—", "—", "—"],
        ["导入 / 归类", "✓", "✓", "—", "—"],
        ["批注 / 复习", "✓", "✓", "—", "—"],
        ["删除课程", "✓", "—", "—", "—"],
        ["管理用户", "—", "—", "—", "✓"],
    ]
    table(slide, 0.72, 1.65, 6.6, 3.1, rows, [2.0, 1.1, 1.1, 1.1, 1.3], font_size=11, header_size=11)
    bullets(slide, 0.72, 5.0, 6.6, 1.6, [
        "页面 / 服务层先判角色，失败返回 403",
        "SQL 再按 sources.owner_id 与 course_members 限制范围",
        "管理员看不到正文、批注与复习内容",
    ], size=12, gap=5)
    picture_fit(slide, "15-permission-denied.png", 7.65, 1.62, 4.95, 4.75)
    text(slide, 7.65, 6.48, 4.95, 0.42,
         "bob 是数据库课程 viewer；直接访问 /courses/1/edit 被 403 拦截。",
         size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
    notes(slide, "35 秒：讲矩阵；现场用 bob 访问编辑 URL 展示 403，admin 只做聚合。")


def slide_tests(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "13 / 测试与验收", "200 项测试与 33 项端到端验收全部通过", 14)

    stats = [("200", "自动化测试"), ("33", "端到端检查"), ("0", "失败项")]
    for index, (value, label) in enumerate(stats):
        x = 0.72 + index * 2.15
        rect(slide, x, 1.62, 1.95, 1.35, WHITE, radius=True, line=LINE, width=0.7)
        text(slide, x + 0.15, 1.72, 1.65, 0.62, value, size=30, bold=True, color=GREEN)
        text(slide, x + 0.15, 2.37, 1.65, 0.3, label, size=10.5, color=MUTED)

    code_block(slide, 0.72, 3.3, 6.0, 2.4,
               ".venv/Scripts/python.exe -m pytest -q\n"
               "# 117 unit + 83 integration\n\n"
               ".venv/Scripts/python.exe scripts/verify.py \\\n"
               "  --json data/verification/report.json", size=10.5)
    text(slide, 0.9, 5.78, 5.6, 0.75,
         "输出：33/33 PASS，0 FAIL，约 0.9 秒\n所有写入检查在结束时整体回滚。", size=11.5, color=GREEN)

    right = [
        "Schema：表、FK、UNIQUE、CHECK、索引、触发器",
        "权限：alice / bob / viewer / admin 边界",
        "导入：幂等、重复、损坏文件、平台适配",
        "检索：关键词、课程、标签、全局搜索",
        "统计：种子数据下的计数与排名",
        "采集：令牌哈希、限流、Bearer 幂等入库",
        "Text2SQL：恶意 SQL、跨用户、只读兜底",
    ]
    for index, item in enumerate(right):
        y = 1.62 + index * 0.68
        rect(slide, 7.1, y, 5.5, 0.54, SOFT if index % 2 == 0 else WHITE,
             radius=True, line=LINE, width=0.6)
        text(slide, 7.28, y + 0.11, 5.15, 0.3, "✓  " + item, size=10.5, color=INK)
    notes(slide, "30 秒：报数字，强调集成测试使用独立 _test 库、验收写入全部回滚。")


def slide_ai(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "14 / AI 协作", "哪些交给 AI，哪些必须自己判断", 15)

    rows = [
        ["被否方案", "最终选择", "原因"],
        ["复用 ReviewMate 向量模型", "经典关系模型", "作业要求实体、联系、SQL"],
        ["会话冗余 owner_id", "经 source 追溯", "避免传递依赖"],
        ["生成列存 content_hash", "md5 表达式索引", "避免 3NF 争议"],
        ["仅按 annotation_id 唯一", "(user_id, annotation_id)", "多人可各自复习"],
        ["只靠 prompt 保护 T2SQL", "只读角色 + 视图", "数据库层兜底"],
        ["中文全文检索", "pg_trgm + ILIKE", "中文分词不可靠"],
    ]
    table(slide, 0.72, 1.62, 8.1, 3.7, rows, [3.25, 2.65, 2.2], font_size=10, header_size=10)

    card(slide, 9.1, 1.62, 3.5, 1.55, "AI 擅长", "生成样板代码、补测试、\n审查 SQL 与文档一致性。", accent=GREEN)
    card(slide, 9.1, 3.35, 3.5, 1.55, "人负责", "业务语义、权限边界、\n取舍理由与最终验收。", accent=RUST)
    rect(slide, 0.72, 5.62, 11.88, 0.95, SOFT, radius=True)
    text(slide, 0.95, 5.75, 11.4, 0.7,
         "协作原则：AI 可以提出方案，但不能替代对数据模型与安全边界的判断；\n"
         "每一次拒绝与修正都记录在 docs/ai-collaboration.md。", size=12.5, color=GREEN)
    notes(slide, "30 秒：举 2 个被否方案，说明人机分工；不把 AI 生成当作正确性证明。")


def slide_end(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, GREEN_DARK)
    text(slide, 0.85, 0.8, 7.6, 0.32, "15 / 结论", size=12, bold=True,
         color=RGBColor(0xD9, 0xC8, 0xA8))
    text(slide, 0.85, 1.45, 8.3, 0.9, "一个可运行、可验证、可复现的\n数据库课程项目", size=34, bold=True,
         color=WHITE, spacing=1.18)

    left = [
        "8 个核心实体，14 张表，完整 E-R 与范式分析",
        "导入、检索、批注、复习、统计、导出的闭环",
        "令牌 + 扩展 + 平台导出，离线也可读可复习",
        "Text2SQL 由只读视图与角色做最终兜底",
        "200 项测试 + 33 项验收，全部通过",
        "README 五条命令从零复现",
    ]
    bullets(slide, 0.85, 3.15, 6.4, 2.4, left, size=14, color=RGBColor(0xE8, 0xE3, 0xD3), gap=8,
           bullet_color=RGBColor(0xD9, 0xC8, 0xA8))

    rect(slide, 7.55, 1.55, 5.0, 4.0, RGBColor(0x2B, 0x50, 0x44), radius=True)
    text(slide, 7.82, 1.78, 4.45, 0.32, "已知限制", size=13, bold=True,
         color=RGBColor(0xD9, 0xC8, 0xA8))
    limitations = [
        "单机课程项目，未做分布式与高并发",
        "420 条消息下规划器通常选顺序扫描",
        "默认离线 mock，配置 Key 后走真实模型",
        "演示数据为模拟问答，不含真实个人信息",
    ]
    bullets(slide, 7.82, 2.3, 4.35, 2.6, limitations, size=11.5,
           color=RGBColor(0xC9, 0xD3, 0xCB), gap=8,
           bullet_color=RGBColor(0xB4, 0x69, 0x3F))

    text(slide, 0.85, 6.35, 6.0, 0.5, "谢谢，欢迎提问", size=20, bold=True, color=WHITE)
    text(slide, 0.85, 6.9, 8.5, 0.3,
         "docs/report.md · docs/design.md · docs/demo-script.md · data/verification/report.json",
         size=9.5, color=RGBColor(0x9F, 0xB3, 0xA6))
    notes(slide, "收尾 15 秒：一句结论 + 主动说明限制，然后把画面交给报告或 Q&A。")


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "个人学习空间：把聊天记录变成可复习的课程资料"
    prs.core_properties.author = "《数据库系统概论》大作业"
    prs.core_properties.subject = "选题 2 · 数据库应用系统"
    prs.core_properties.keywords = "PostgreSQL, FastAPI, Text2SQL, E-R, 数据库设计"

    title_slide(prs)
    slide_why(prs)
    slide_scope(prs)
    slide_data(prs)
    slide_er(prs)
    slide_schema(prs)
    slide_norm(prs)
    slide_import(prs)
    slide_search(prs)
    slide_review(prs)
    slide_stats(prs)
    slide_t2s(prs)
    slide_permissions(prs)
    slide_tests(prs)
    slide_ai(prs)
    slide_end(prs)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"已生成 {path.relative_to(ROOT)}（16 页）")
